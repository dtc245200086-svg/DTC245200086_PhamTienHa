from __future__ import annotations

import html
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

REPORT_DIR = Path(__file__).resolve().parent
SOURCE = REPORT_DIR / "Final_Report.md"


def split_cells(line: str) -> list[str]:
    return [cell.strip().replace(r"\|", "|") for cell in re.split(r"(?<!\\)\|", line.strip().strip("|"))]


def inline_html(value: str) -> str:
    value = html.escape(value)
    value = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", value)
    value = re.sub(r"`([^`]+)`", r"<code>\1</code>", value)
    value = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r'<a href="\2">\1</a>', value)
    return value


def is_separator(row: list[str]) -> bool:
    return all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in row)


def render_page(markdown: str) -> str:
    lines = markdown.splitlines()
    output: list[str] = []
    index = 0
    in_list: str | None = None

    def close_list() -> None:
        nonlocal in_list
        if in_list:
            output.append(f"</{in_list}>")
            in_list = None

    while index < len(lines):
        line = lines[index].rstrip()
        if not line.strip():
            close_list()
            index += 1
            continue

        image = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line.strip())
        if image:
            close_list()
            alt = html.escape(image.group(1))
            src = html.escape(image.group(2), quote=True)
            output.append(f'<figure><img src="{src}" alt="{alt}"><figcaption>{alt}</figcaption></figure>')
            index += 1
            continue

        if line.strip().startswith("|"):
            close_list()
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_cells(lines[index]))
                index += 1
            rows = [row for row in rows if not is_separator(row)]
            if rows:
                output.append("<table><thead><tr>" + "".join(f"<th>{inline_html(cell)}</th>" for cell in rows[0]) + "</tr></thead><tbody>")
                for row in rows[1:]:
                    output.append("<tr>" + "".join(f"<td>{inline_html(cell)}</td>" for cell in row) + "</tr>")
                output.append("</tbody></table>")
            continue

        if line.startswith("```"):
            close_list()
            index += 1
            code: list[str] = []
            while index < len(lines) and not lines[index].startswith("```"):
                code.append(lines[index])
                index += 1
            index += 1
            output.append("<pre><code>" + html.escape("\n".join(code)) + "</code></pre>")
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)", line)
        if heading:
            close_list()
            level = len(heading.group(1))
            output.append(f"<h{level}>{inline_html(heading.group(2))}</h{level}>")
            index += 1
            continue

        if line.strip() == "---":
            close_list()
            output.append("<hr>")
            index += 1
            continue

        bullet = re.match(r"^\s*[-*]\s+(.*)", line)
        numbered = re.match(r"^\s*\d+\.\s+(.*)", line)
        if bullet or numbered:
            tag = "ul" if bullet else "ol"
            value = bullet.group(1) if bullet else numbered.group(1)
            if in_list != tag:
                close_list()
                output.append(f"<{tag}>")
                in_list = tag
            output.append(f"<li>{inline_html(value)}</li>")
            index += 1
            continue

        close_list()
        paragraph = [line.strip()]
        index += 1
        while index < len(lines):
            next_line = lines[index]
            if not next_line.strip() or next_line.lstrip().startswith(("#", "|", "```", "- ", "* ")) or re.match(r"^\s*\d+\.\s+", next_line):
                break
            paragraph.append(next_line.strip())
            index += 1
        output.append("<p>" + inline_html(" ".join(paragraph)) + "</p>")

    close_list()
    return "\n".join(output)


def add_runs(paragraph, value: str) -> None:
    tokens = re.split(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))", value)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
        elif token.startswith("["):
            match = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            paragraph.add_run(match.group(1) if match else token)
        else:
            paragraph.add_run(token)


def add_docx_page(document: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    index = 0
    in_code = False
    code: list[str] = []
    while index < len(lines):
        line = lines[index].rstrip()
        if line.startswith("```"):
            if not in_code:
                in_code = True
                code = []
            else:
                paragraph = document.add_paragraph()
                run = paragraph.add_run("\n".join(code))
                run.font.name = "Consolas"
                run.font.size = Pt(8)
                in_code = False
            index += 1
            continue
        if in_code:
            code.append(line)
            index += 1
            continue

        image = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line.strip())
        if image:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            image_path = (REPORT_DIR / image.group(2)).resolve()
            if image_path.exists():
                paragraph.add_run().add_picture(str(image_path), width=Inches(6.1))
            caption = document.add_paragraph(image.group(1))
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in caption.runs:
                run.italic = True
                run.font.size = Pt(8)
            index += 1
            continue

        heading = re.match(r"^(#{1,3})\s+(.*)", line)
        if heading:
            paragraph = document.add_heading(heading.group(2), level=len(heading.group(1)))
            if heading.group(1) == "#":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            index += 1
            continue

        if line.strip().startswith("|"):
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_cells(lines[index]))
                index += 1
            rows = [row for row in rows if not is_separator(row)]
            if rows:
                table = document.add_table(rows=1, cols=len(rows[0]))
                table.style = "Light Shading Accent 1"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for column, value in enumerate(rows[0]):
                    add_runs(table.rows[0].cells[column].paragraphs[0], value)
                for row in rows[1:]:
                    cells = table.add_row().cells
                    for column, value in enumerate(row[:len(cells)]):
                        add_runs(cells[column].paragraphs[0], value)
            continue

        bullet = re.match(r"^\s*[-*]\s+(.*)", line)
        numbered = re.match(r"^\s*\d+\.\s+(.*)", line)
        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            add_runs(paragraph, bullet.group(1))
        elif numbered:
            paragraph = document.add_paragraph(style="List Number")
            add_runs(paragraph, numbered.group(1))
        elif line.strip() and line.strip() != "---":
            paragraph = document.add_paragraph()
            add_runs(paragraph, line)
        index += 1


def build_html(pages: list[str]) -> None:
    css = """@page{size:A4;margin:15mm 16mm 17mm}*{box-sizing:border-box}body{font-family:Arial,sans-serif;color:#202833;font-size:9.5pt;line-height:1.28;margin:0}.page{page-break-after:always;break-after:page;min-height:265mm;position:relative}.page:last-child{page-break-after:auto;break-after:auto}h1{font-size:19pt;color:#17324d;margin:0 0 12pt}h2{font-size:15pt;color:#245477;margin:0 0 12pt}h3{font-size:11pt;color:#245477;margin:10pt 0 4pt}p{margin:0 0 7pt}ul,ol{margin:3pt 0 7pt;padding-left:20pt}li{margin:0 0 3pt}table{border-collapse:collapse;width:100%;font-size:8pt;margin:6pt 0 9pt;page-break-inside:avoid}th,td{border:1px solid #97a9b8;padding:4pt 5pt;vertical-align:top}th{background:#e8f0f5;color:#17324d;text-align:left}pre{white-space:pre-wrap;background:#f2f5f7;border:1px solid #d7e0e6;padding:7pt;font-family:Consolas,monospace;font-size:8pt}code{font-family:Consolas,monospace;font-size:.94em}figure{margin:6pt 0 8pt;text-align:center;page-break-inside:avoid}figure img{max-width:100%;max-height:105mm;object-fit:contain}figcaption{font-size:8pt;color:#526575;margin-top:3pt}a{color:#145d87;text-decoration:none}hr{border:0;border-top:1px solid #cad4dc;margin:8pt 0}"""
    content = "".join(f'<section class="page">{render_page(page)}</section>' for page in pages)
    (REPORT_DIR / "Final_Report.html").write_text(f'<!doctype html><html lang="vi"><head><meta charset="utf-8"><title>Báo cáo Đề 18 Billing</title><style>{css}</style></head><body>{content}</body></html>', encoding="utf-8")


def build_docx(pages: list[str]) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(9.5)
    document.styles["Normal"].paragraph_format.space_after = Pt(5)
    for name, size in (("Title", 22), ("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 11)):
        document.styles[name].font.name = "Arial"
        document.styles[name].font.size = Pt(size)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Đề 18 — Billing | ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for index, page in enumerate(pages):
        add_docx_page(document, page)
        if index < len(pages) - 1:
            document.add_page_break()
    document.save(REPORT_DIR / "Final_Report.docx")


def main() -> None:
    source = SOURCE.read_text(encoding="utf-8")
    pages = [page.strip() for page in source.split("<!-- PAGE BREAK -->") if page.strip()]
    build_html(pages)
    build_docx(pages)
    print(f"REPORT_SECTIONS={len(pages)}")
    print("HTML_AND_DOCX=PASS")


if __name__ == "__main__":
    main()
