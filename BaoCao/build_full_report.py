# -*- coding: utf-8 -*-
"""
Full Report Generator for Billing_One (Topic 18 - De 18: He thong Quan ly Hoa don / Billing)
Generates:
1. BAO_CAO_TRIEN_KHAI_HE_THONG_BILLING_ONE.docx (Formatted according to standard academic form)
2. BAO_CAO_TRIEN_KHAI_HE_THONG_BILLING_ONE.md
3. BAO_CAO_TRIEN_KHAI_HE_THONG_BILLING_ONE.html
4. BAO_CAO_TRIEN_KHAI_HE_THONG_BILLING_ONE.pdf (via Edge headless)
"""

import os
import sys
import subprocess
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = Path(r"c:\Users\phamt\OneDrive\Desktop\Billing_One")
BAOCAO_DIR = BASE_DIR / "BaoCao"
EVIDENCE_DIR = BASE_DIR / "docs" / "evidence"

BAOCAO_DIR.mkdir(parents=True, exist_ok=True)

# Helper XML styling functions for Word
def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="B0C4DE", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_callout_box(doc, text, title="LƯU Ý / GHI CHÚ QUAN TRỌNG:"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    # left thick border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="17324D"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(f"{title}\n")
    r1.bold = True
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(11)
    r1.font.color.rgb = RGBColor(23, 50, 77)
    r2 = p.add_run(text)
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(11)
    r2.font.color.rgb = RGBColor(40, 40, 40)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

print("Helper functions defined successfully.")
